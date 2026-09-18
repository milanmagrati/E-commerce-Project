"""Verify the bulk-log batch export (Excel + CSV, every scope, both providers).

Run:  python test_bulk_log_export.py

Exercises the real views through Django's test client against the live DB, so
it also proves the URLs, the permission decorator and the download headers are
wired up - not just the row-building helpers.
"""
import os
import csv
import io

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'myproject.settings')
django.setup()

from django.conf import settings  # noqa: E402

# The test client speaks to host "testserver", which the project's real
# ALLOWED_HOSTS has no reason to contain.
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS = list(settings.ALLOWED_HOSTS) + ['testserver']

from django.db import transaction  # noqa: E402
from django.test import Client  # noqa: E402
from django.test.utils import setup_test_environment  # noqa: E402
from django.urls import reverse  # noqa: E402

# response.context is only captured when the instrumented template renderer is
# installed; a standalone script has to ask for it explicitly.
try:
    setup_test_environment()
except RuntimeError:
    pass
from openpyxl import load_workbook  # noqa: E402

from accounts.models import CustomUser  # noqa: E402
from ncm.models import NCMBulkLog, NCMBulkLogOrder  # noqa: E402
from pick_and_drop.models import PNDBulkLog, PNDBulkLogOrder  # noqa: E402
from dashboard.views import (  # noqa: E402
    BULK_LOG_EXPORT_SCOPES,
    _bulk_log_excel_value,
    _bulk_log_export_columns,
    _bulk_log_export_counts,
    _bulk_log_export_rows,
)

failures = []


def check(label, condition, detail=''):
    status = 'PASS' if condition else 'FAIL'
    print('  [{}] {}{}'.format(status, label, ' - ' + detail if detail else ''))
    if not condition:
        failures.append(label)


def pick_batch(log_model, order_model):
    """Prefer a batch that has non-success rows - that is where scoping bites."""
    fallback = None
    for candidate in log_model.objects.filter(is_deleted=False).order_by('-created_at')[:50]:
        rows = order_model.objects.filter(batch=candidate)
        if not rows.exists():
            continue
        if fallback is None:
            fallback = candidate
        if rows.exclude(status='success').exists():
            return candidate
    return fallback


def run_provider(label, log_model, order_model, id_label, id_field, url_name, admin):
    print('\n=== {} ==='.format(label))
    bulk_log = pick_batch(log_model, order_model)
    if bulk_log is None:
        print('  [SKIP] no {} batch with orders in this database'.format(label))
        return

    batch_orders = order_model.objects.filter(batch=bulk_log).select_related('order__api_config')
    columns = _bulk_log_export_columns(id_label)
    counts = _bulk_log_export_counts(batch_orders)
    print('  Batch under test: {} ({} rows: {} ok / {} failed / {} skipped)'.format(
        bulk_log.batch_number, counts['all'], counts['success'],
        counts['failed'], counts['skipped']))

    # -- 1. Row builder: scopes partition the batch, columns line up ----------
    print('  1. Row builder')
    rows = {scope: _bulk_log_export_rows(batch_orders, scope, id_field)
            for scope in BULK_LOG_EXPORT_SCOPES}

    check('all scope covers every recorded row',
          len(rows['all']) == batch_orders.count(),
          '{} rows vs {} entries'.format(len(rows['all']), batch_orders.count()))
    check('the three status scopes partition the batch',
          len(rows['success']) + len(rows['failed']) + len(rows['skipped']) == len(rows['all']))
    for scope in ('success', 'failed', 'skipped'):
        check('{} scope matches its row count'.format(scope),
              len(rows[scope]) == counts[scope],
              '{} vs {}'.format(len(rows[scope]), counts[scope]))
    check('every row has one cell per column',
          all(len(r) == len(columns) for r in rows['all']))
    check('column header carries the provider tracking id',
          id_label in columns, id_label)

    status_idx = columns.index('Status')
    for scope, display in (('success', 'Success'), ('failed', 'Failed'), ('skipped', 'Skipped')):
        check('{} scope holds only {} rows'.format(scope, display),
              all(r[status_idx] == display for r in rows[scope]))
    check('row numbering restarts at 1 per scope',
          all(not rs or rs[0][0] == 1 for rs in rows.values()))

    amount_idx = columns.index('COD Amount')
    check('COD amounts export as numbers, not strings',
          all(isinstance(r[amount_idx], float) for r in rows['all']))

    # -- 2. The views themselves, through the URL, as a logged-in admin ------
    print('  2. HTTP responses')
    if admin is None:
        print('  [SKIP] no administrator user in this database')
        return

    client = Client()
    client.force_login(admin)
    url = reverse(url_name, args=[bulk_log.id])

    xlsx = client.get(url, {'format': 'xlsx', 'scope': 'all'})
    check('xlsx returns 200', xlsx.status_code == 200, str(xlsx.status_code))
    check('xlsx is an attachment named after the batch',
          xlsx['Content-Disposition'].startswith('attachment')
          and bulk_log.batch_number in xlsx['Content-Disposition'],
          xlsx.get('Content-Disposition', ''))
    check('xlsx has the spreadsheet content type',
          'spreadsheetml' in xlsx['Content-Type'])

    wb = load_workbook(io.BytesIO(xlsx.content))
    check('workbook has Summary + Orders sheets',
          wb.sheetnames == ['Summary', 'Orders'], str(wb.sheetnames))
    summary_labels = [row[0] for row in wb['Summary'].iter_rows(values_only=True)]
    check('summary names the batch and the exported scope',
          'Batch Number' in summary_labels and 'Rows Exported' in summary_labels,
          str(summary_labels))
    sheet = wb['Orders']
    check('sheet header matches the shared column list',
          [c.value for c in sheet[1]] == columns)
    check('sheet has one data row per batch entry',
          sheet.max_row - 1 == len(rows['all']),
          '{} data rows vs {} entries'.format(sheet.max_row - 1, len(rows['all'])))
    check('header row is frozen', sheet.freeze_panes == 'A2')

    scope = 'failed' if counts['failed'] else 'success'
    csv_resp = client.get(url, {'format': 'csv', 'scope': scope})
    check('csv returns 200', csv_resp.status_code == 200, str(csv_resp.status_code))
    check('csv filename reflects the scope',
          BULK_LOG_EXPORT_SCOPES[scope] + '.csv' in csv_resp['Content-Disposition'],
          csv_resp.get('Content-Disposition', ''))
    check('csv is BOM-prefixed for Excel', csv_resp.content.startswith(b'\xef\xbb\xbf'))
    parsed = list(csv.reader(io.StringIO(csv_resp.content.decode('utf-8-sig'))))
    check('csv header matches the shared column list', parsed[0] == columns)
    check('csv body matches the {} scope'.format(scope),
          len(parsed) - 1 == len(rows[scope]),
          '{} rows vs {}'.format(len(parsed) - 1, len(rows[scope])))

    bad_scope = client.get(url, {'format': 'xlsx', 'scope': 'nonsense'})
    check('unknown scope falls back to all instead of erroring',
          bad_scope.status_code == 200
          and 'all-orders' in bad_scope['Content-Disposition'],
          bad_scope.get('Content-Disposition', ''))

    default_fmt = client.get(url)
    check('no format defaults to xlsx',
          default_fmt.status_code == 200
          and 'spreadsheetml' in default_fmt['Content-Type'])

    missing = client.get(reverse(url_name, args=[999999999]))
    check('unknown batch id is a 404, not a crash', missing.status_code == 404,
          str(missing.status_code))

    anon = Client().get(url)
    check('anonymous request is not served the file',
          anon.status_code in (302, 401, 403), str(anon.status_code))

    # -- 3. The detail page offers the export -------------------------------
    print('  3. Detail page wiring')
    detail_name = url_name.replace('_export', '_detail')
    page = client.get(reverse(detail_name, args=[bulk_log.id]))
    check('detail page renders', page.status_code == 200, str(page.status_code))
    body = page.content.decode('utf-8', 'ignore')
    check('detail page links the export', url in body)
    check('detail page counts match the exportable rows',
          page.context['export_counts'] == counts,
          '{} vs {}'.format(page.context.get('export_counts'), counts))


# -- Cell sanitising is provider-independent, so check it once ---------------
print('\n=== Excel cell sanitising ===')
check('control characters are stripped', _bulk_log_excel_value('ok\x01bad') == 'okbad')
check('over-long text is truncated to a writable length',
      len(_bulk_log_excel_value('x' * 40000)) == 32000)
check('non-strings pass through untouched', _bulk_log_excel_value(12.5) == 12.5)

class _Rollback(Exception):
    """Thrown to unwind the synthetic batch - nothing below is persisted."""


def run_synthetic(label, log_model, order_model, id_label, id_field, url_name,
                  admin, log_kwargs, tracking_value):
    """Cover the rows the live DB happens not to have: failed, skipped,
    Nepali text, a missing tracking id and a control character in the message.

    Everything is written inside a transaction that is rolled back, so the
    developer's database is left exactly as it was found.
    """
    print('\n=== {} (synthetic batch) ==='.format(label))
    if admin is None:
        print('  [SKIP] no administrator user in this database')
        return

    columns = _bulk_log_export_columns(id_label)
    client = Client()
    client.force_login(admin)

    try:
        with transaction.atomic():
            bulk_log = log_model.objects.create(
                batch_number=log_model.generate_batch_number(),
                total_orders=3, success_count=1, failed_count=1, skipped_count=1,
                status='partial', created_by=admin, **log_kwargs)

            entries = [
                ('success', tracking_value, 'Sent to courier'),
                ('failed', None, 'Rejected\x01by API ' + 'x' * 40000),
                ('skipped', None, 'Already sent'),
            ]
            for i, (status, tracking, message) in enumerate(entries, start=1):
                order_model.objects.create(
                    batch=bulk_log,
                    order_number='SYN{}'.format(i),
                    customer_name='सबिता श्रेष्ठ',
                    customer_phone='980000000{}'.format(i),
                    shipping_address='काठमाडौं, नयाँ बसपार्क',
                    cod_amount=1500,
                    destination_branch='TINKUNE',
                    status=status,
                    message=message,
                    **{id_field: tracking})

            batch_orders = order_model.objects.filter(batch=bulk_log)
            counts = _bulk_log_export_counts(batch_orders)
            check('synthetic counts split 1/1/1',
                  counts == {'all': 3, 'success': 1, 'failed': 1, 'skipped': 1},
                  str(counts))

            url = reverse(url_name, args=[bulk_log.id])
            xlsx = client.get(url, {'format': 'xlsx', 'scope': 'all'})
            check('mixed-status xlsx returns 200', xlsx.status_code == 200,
                  str(xlsx.status_code))

            sheet = load_workbook(io.BytesIO(xlsx.content))['Orders']
            check('all three rows are written', sheet.max_row - 1 == 3,
                  str(sheet.max_row - 1))
            fills = {sheet.cell(row=r, column=columns.index('Status') + 1).value:
                     sheet.cell(row=r, column=1).fill.start_color.rgb
                     for r in range(2, sheet.max_row + 1)}
            check('failed rows are tinted red', str(fills.get('Failed')).endswith('FFE4E6'),
                  str(fills.get('Failed')))
            check('skipped rows are tinted amber', str(fills.get('Skipped')).endswith('FEF3C7'),
                  str(fills.get('Skipped')))
            check('success rows are left untinted',
                  not str(fills.get('Success')).endswith(('FFE4E6', 'FEF3C7')),
                  str(fills.get('Success')))

            message_col = columns.index('Message') + 1
            written = [sheet.cell(row=r, column=message_col).value
                       for r in range(2, sheet.max_row + 1)]
            check('the control character never reaches the workbook',
                  all('\x01' not in (v or '') for v in written))
            check('the over-long message is truncated, not dropped',
                  any(v and v.startswith('Rejectedby API') and len(v) == 32000
                      for v in written),
                  str([len(v or '') for v in written]))

            id_col = columns.index(id_label) + 1
            check('a missing tracking id exports blank, not "None"',
                  sheet.cell(row=3, column=id_col).value in (None, ''),
                  repr(sheet.cell(row=3, column=id_col).value))

            failed_csv = client.get(url, {'format': 'csv', 'scope': 'failed'})
            body = failed_csv.content.decode('utf-8-sig')
            parsed = list(csv.reader(io.StringIO(body)))
            check('failed-only csv carries exactly the failed row',
                  len(parsed) - 1 == 1, str(len(parsed) - 1))
            check('Nepali text survives the csv round-trip',
                  'सबिता श्रेष्ठ' in body)

            skipped = client.get(url, {'format': 'xlsx', 'scope': 'skipped'})
            check('skipped scope names itself in the filename',
                  'skipped.xlsx' in skipped['Content-Disposition'],
                  skipped.get('Content-Disposition', ''))

            raise _Rollback
    except _Rollback:
        pass

    check('synthetic batch left no trace',
          not log_model.objects.filter(created_by=admin, total_orders=3,
                                       success_count=1, failed_count=1,
                                       skipped_count=1, status='partial')
          .filter(orders__order_number='SYN1').exists())


admin = (CustomUser.objects.filter(is_superuser=True).first()
         or CustomUser.objects.filter(role='administrator').first())

run_provider('NCM', NCMBulkLog, NCMBulkLogOrder, 'NCM ID', 'ncm_order_id',
             'ncm_bulk_log_export', admin)
run_provider('Pick and Drop', PNDBulkLog, PNDBulkLogOrder, 'PND ID', 'pnd_order_id',
             'pnd_bulk_log_export', admin)

run_synthetic('NCM', NCMBulkLog, NCMBulkLogOrder, 'NCM ID', 'ncm_order_id',
              'ncm_bulk_log_export', admin,
              {'from_branch': 'TINKUNE', 'delivery_type': 'Door2Door'}, 25137563)
run_synthetic('Pick and Drop', PNDBulkLog, PNDBulkLogOrder, 'PND ID', 'pnd_order_id',
              'pnd_bulk_log_export', admin,
              {'destination_branch': 'KATHMANDU VALLEY'}, 'PND-99001')

print('\n' + ('ALL CHECKS PASSED' if not failures else 'FAILED: ' + ', '.join(failures)))
raise SystemExit(1 if failures else 0)
