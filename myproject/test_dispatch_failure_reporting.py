"""
Verify dispatch Success/Failed reporting is accurate end to end.

Covers the defects this change fixes:
  1. A scanned ID that matched no order was counted as neither Success nor
     Failed, so the Failed column read 0 for a batch that wholly failed.
  2. Not-found scans carried no failure_reason, so the detail page could not
     say why anything failed.
  3. A failed prior attempt permanently blocked re-dispatching an order.
  4. Rows with failures were visually identical to clean rows on the list.

Run:  python test_dispatch_failure_reporting.py
"""
import os
import sys
from html.parser import HTMLParser

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'myproject.settings')
django.setup()

from django.conf import settings  # noqa: E402

# The test client sends Host: testserver; this script talks to the real
# configured database, so widen only the host check.
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS = list(settings.ALLOWED_HOSTS) + ['testserver']

from django.db import transaction  # noqa: E402
from django.test import Client  # noqa: E402
from django.urls import reverse  # noqa: E402

from accounts.models import CustomUser  # noqa: E402
from dashboard.models import (  # noqa: E402
    Dispatch, DispatchItem, DispatchLog, Order,
)

FAILURES = []


class _CardCellAudit(HTMLParser):
    """Below 768px the dispatch tables collapse into cards, and each cell shows
    its column name from `data-label`. A cell without one renders as a bare
    value with no indication of what it means, so assert full coverage."""

    def __init__(self):
        super().__init__()
        self.in_card_table = False
        self.row_class = ''
        self.total = 0
        self.unlabelled = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        css = attrs.get('class', '')
        if tag == 'table':
            self.in_card_table = 'rcards' in css or 'rtable' in css
        elif tag == 'tr':
            self.row_class = css
        elif tag == 'td' and self.in_card_table and 'no-card' not in self.row_class:
            self.total += 1
            if 'data-label' not in attrs:
                self.unlabelled += 1

    def handle_endtag(self, tag):
        if tag == 'table':
            self.in_card_table = False


def check(label, condition, detail=''):
    status = 'PASS' if condition else 'FAIL'
    print(f'  [{status}] {label}' + (f'  -> {detail}' if detail and not condition else ''))
    if not condition:
        FAILURES.append(label)


def get_staff_user():
    user = CustomUser.objects.filter(is_superuser=True).first()
    if not user:
        raise SystemExit('No superuser exists; create one before running this script.')
    return user


def main():
    user = get_staff_user()
    print(f'Using account: {user.username}\n')

    # ── Fixtures: one real order, one ID that matches nothing ────────────────
    order = Order.objects.filter(is_deleted=False).exclude(order_status='dispatched').first()
    if not order:
        raise SystemExit('No non-dispatched order available to exercise the dispatch flow.')

    bogus_id = 'ZZ-NO-SUCH-ORDER-0001'
    real_id = order.order_number

    print('--- 1. Counts include not-found scans -------------------------------')
    with transaction.atomic():
        batch = Dispatch.objects.create(
            batch_number='TEST-DISPATCH-COUNTS',
            logistics='ncm', status='dispatched', total_orders=3, created_by=user,
        )
        ok = DispatchItem.objects.create(dispatch=batch, scanned_order_id=real_id, order=order)
        ok.mark_success()
        rejected = DispatchItem.objects.create(dispatch=batch, scanned_order_id='DUP-1', order=order)
        rejected.mark_failed('Already dispatched in batch TEST-EARLIER', code='already_dispatched')
        missing = DispatchItem.objects.create(dispatch=batch, scanned_order_id=bogus_id)
        missing.mark_not_found()

        batch.refresh_from_db()
        check('success count is 1', batch.get_success_count() == 1, batch.get_success_count())
        check('failed count is 1', batch.get_failed_count() == 1, batch.get_failed_count())
        check('not-found count is 1', batch.get_not_found_count() == 1, batch.get_not_found_count())
        check('Failed column (issues) counts BOTH failure kinds = 2',
              batch.get_issue_count() == 2, batch.get_issue_count())
        check('outcome is partial', batch.get_outcome() == 'partial', batch.get_outcome())
        check('success rate is 33%', batch.get_success_rate() == 33, batch.get_success_rate())

        print('\n--- 2. Every failed scan carries a reason --------------------------')
        for item in batch.items.all():
            if item.dispatch_status == 'success':
                continue
            check(f'{item.scanned_order_id} has a reason',
                  bool(item.get_failure_reason_display()))
            check(f'{item.scanned_order_id} has a failure_code',
                  bool(item.failure_code), item.failure_code)
            check(f'{item.scanned_order_id} has failed_at set', item.failed_at is not None)
        check('not-found scan is coded not_found',
              batch.items.get(scanned_order_id=bogus_id).failure_code == 'not_found')

        print('\n--- 3. Legacy batch with no item rows is not silently 0/0 ----------')
        legacy = Dispatch.objects.create(
            batch_number='TEST-DISPATCH-LEGACY',
            logistics='ncm', status='dispatched', total_orders=4, created_by=user,
        )
        check('legacy outcome is "unrecorded"', legacy.get_outcome() == 'unrecorded', legacy.get_outcome())
        check('legacy reports 4 unrecorded scans', legacy.get_unrecorded_count() == 4,
              legacy.get_unrecorded_count())

        print('\n--- 4. Audit log records batch events ------------------------------')
        batch.log('order_failed', 'test event', level='error', item=rejected, user=user)
        check('log row written', DispatchLog.objects.filter(dispatch=batch).count() == 1)

        print('\n--- 5. List page renders accurate columns + row colour -------------')
        client = Client()
        client.force_login(user)

        resp = client.get(reverse('dispatch_list'))
        check('list page returns 200', resp.status_code == 200, resp.status_code)
        html = resp.content.decode('utf-8', 'replace')
        check('failing batch row is colour-coded', 'row-partial' in html or 'row-failed' in html)
        check('list exposes an outcome filter', 'name="outcome"' in html)
        check('list warns about undispatched scans', 'did not dispatch' in html)

        resp = client.get(reverse('dispatch_list') + '?outcome=issues')
        check('outcome=issues filter returns 200', resp.status_code == 200, resp.status_code)
        listed = resp.content.decode('utf-8', 'replace')
        check('outcome=issues includes the failing batch', 'TEST-DISPATCH-COUNTS' in listed)
        check('outcome=issues excludes the batch with no failures',
              'TEST-DISPATCH-LEGACY' not in listed)

        print('\n--- 6. Detail page shows the reason for each failure ----------------')
        resp = client.get(reverse('dispatch_detail', args=[batch.pk]))
        check('detail page returns 200', resp.status_code == 200, resp.status_code)
        html = resp.content.decode('utf-8', 'replace')
        check('failure report is present', 'Failure Report' in html)
        check('rejection reason is shown', 'Already dispatched in batch TEST-EARLIER' in html)
        check('not-found reason is shown', 'Order ID not found in the system' in html)
        check('remediation hint is shown', 'mis-scan' in html)
        check('activity log is present', 'Activity Log' in html)
        check('failed IDs are offered for copying', 'Copy failed IDs' in html)
        check('per-order table has a Reason column', '<th>Reason</th>' in html)
        check('stock reduction table excludes failed scans', bogus_id not in
              html.split('Stock Reduction Details')[-1])

        print('\n--- 7. A failed attempt does not permanently block re-dispatch ------')
        blocking = DispatchItem.objects.filter(
            order=order, dispatch__is_deleted=False, dispatch_status='success',
        ).exclude(dispatch=batch).exists()
        check('only successful prior scans block a re-dispatch', not blocking)

        transaction.set_rollback(True)

    # ── Live scan flow, exercised through the real POST view ─────────────────
    print('\n--- 8. Real scan flow records outcomes + reasons --------------------')
    with transaction.atomic():
        client = Client()
        client.force_login(user)

        client.post(reverse('dispatch_management'), {
            'order_ids': f'{real_id}, {bogus_id}',
            'set_status': 'dispatched',
            'logistics': 'ncm',
        }, follow=True)
        live = Dispatch.objects.order_by('-id').first()
        check('scan created a batch', live is not None and live.total_orders == 2)
        check('good ID recorded as success', live.get_success_count() == 1, live.get_success_count())
        check('bogus ID recorded as not_found', live.get_not_found_count() == 1,
              live.get_not_found_count())
        check('Failed column would show 1', live.get_issue_count() == 1, live.get_issue_count())
        missing = live.items.get(scanned_order_id=bogus_id)
        check('bogus ID has a stored reason', bool(missing.failure_reason), missing.failure_reason)
        check('batch wrote an audit trail', live.logs.count() >= 3, live.logs.count())
        check('audit trail has a closing summary',
              live.logs.filter(event='batch_completed').exists())

        print('\n--- 9. Re-scanning a dispatched order is rejected with the cause ----')
        resp = client.post(reverse('dispatch_management'), {
            'order_ids': real_id,
            'set_status': 'dispatched',
            'logistics': 'ncm',
        }, follow=True)
        repeat = Dispatch.objects.order_by('-id').first()
        # Batch numbers are second-granular; a same-second resubmit must not
        # collide on the unique constraint and lose the whole batch.
        check('second batch got a distinct batch_number',
              repeat.pk != live.pk and repeat.batch_number != live.batch_number,
              repeat.batch_number)
        check('repeat scan is marked failed', repeat.get_failed_count() == 1,
              repeat.get_failed_count())
        dupe = repeat.items.first()
        check('rejection is coded already_dispatched',
              dupe.failure_code == 'already_dispatched', dupe.failure_code)
        check('reason names the earlier batch', live.batch_number in dupe.failure_reason,
              dupe.failure_reason)
        html = resp.content.decode('utf-8', 'replace')
        check('detail page explains how to fix it', 'Remove it from this batch' in html)

        transaction.set_rollback(True)

    # ── Responsive markup + remaining UI wiring ──────────────────────────────
    print('\n--- 10. Mobile card layout is wired correctly ----------------------')
    client = Client()
    client.force_login(user)

    pages = {'list': reverse('dispatch_list')}
    sample = Dispatch.objects.filter(is_deleted=False, items__isnull=False).first()
    if sample:
        pages['detail'] = reverse('dispatch_detail', args=[sample.pk])

    for name, url in pages.items():
        html = client.get(url).content.decode('utf-8', 'replace')
        audit = _CardCellAudit()
        audit.feed(html)
        check(f'{name}: every card cell carries a data-label',
              audit.unlabelled == 0, f'{audit.unlabelled} of {audit.total}')
        check(f'{name}: has card-mode tables', audit.total > 0, audit.total)

    html = client.get(reverse('dispatch_list')).content.decode('utf-8', 'replace')
    check('list: filter links are not malformed', '?&outcome' not in html
          and '?&amp;outcome' not in html)
    check('list: delete button carries a server-resolved URL', 'data-url="/dispatch/' in html)
    check('list: stat strip present', 'stat-strip' in html)

    print('\n--- 11. A trashed batch offers Restore, not Move to Trash ----------')
    trashed = Dispatch.objects.filter(is_deleted=True).first()
    if trashed:
        html = client.get(reverse('dispatch_detail', args=[trashed.pk])).content.decode('utf-8', 'replace')
        check('trashed batch shows the trash banner', 'in the trash' in html)
        check('trashed batch offers Restore',
              reverse('dispatch_restore', args=[trashed.pk]) in html)
        check('trashed batch hides Move to Trash (it would 404)',
              reverse('dispatch_move_to_trash', args=[trashed.pk]) not in html)
    else:
        print('  [SKIP] no trashed dispatch in this database')

    live = Dispatch.objects.filter(is_deleted=False).first()
    html = client.get(reverse('dispatch_detail', args=[live.pk])).content.decode('utf-8', 'replace')
    check('live batch offers Move to Trash',
          reverse('dispatch_move_to_trash', args=[live.pk]) in html)
    check('live batch hides Restore',
          reverse('dispatch_restore', args=[live.pk]) not in html)

    print('\n--- 12. Excel export carries the failure reasons -------------------')
    ids = list(Dispatch.objects.filter(is_deleted=False).values_list('id', flat=True)[:5])
    resp = client.post(reverse('dispatch_bulk_action'),
                       {'dispatch_ids': ids, 'bulk_action': 'export_excel'})
    check('export returns a workbook, not a redirect', resp.status_code == 200, resp.status_code)
    check('export is served as a download', 'attachment;' in resp.get('Content-Disposition', ''))
    if resp.status_code == 200:
        import io
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(resp.content))
        check('export has both sheets', wb.sheetnames == ['Batches', 'Scanned Orders'], wb.sheetnames)
        headers = [c.value for c in wb['Scanned Orders'][1]]
        check('export includes a Reason column', 'Reason' in headers, headers)

    print('\n' + '=' * 68)
    if FAILURES:
        print(f'{len(FAILURES)} CHECK(S) FAILED:')
        for name in FAILURES:
            print(f'  - {name}')
        sys.exit(1)
    print('All checks passed. (Test data rolled back — database unchanged.)')


if __name__ == '__main__':
    main()
