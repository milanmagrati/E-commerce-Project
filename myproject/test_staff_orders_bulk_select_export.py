"""Verify bulk-selection export on the Staff Performance page.

Covers the two endpoints backing the checkbox "select N, export exactly
those" workflow added on top of the plain filtered export:

  - ?orders_ids_only=1            -> JSON list of ids matching current filters
  - ?orders_export=xlsx&order_ids=1,2,3 -> workbook of exactly those orders

    python test_staff_orders_bulk_select_export.py
"""
import os
import sys
import io

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'myproject.settings')
django.setup()

from django.contrib.auth import get_user_model
from django.contrib.sessions.backends.db import SessionStore
from django.db.models import Min, Max
from django.test import RequestFactory
from django.utils import timezone
from openpyxl import load_workbook
import json

from dashboard.models import Order
from dashboard.views import staff_performance_analytics

User = get_user_model()

PASS, FAIL = [], []


def check(label, condition, detail=''):
    (PASS if condition else FAIL).append(label)
    print(('  [PASS] ' if condition else '  [FAIL] ') + label + (f' -> {detail}' if detail else ''))


def call(user, **params):
    request = RequestFactory().get('/staff-performance/', params)
    request.user = user
    request.session = SessionStore()
    return staff_performance_analytics(request)


def main():
    admin = User.objects.filter(is_superuser=True, is_active=True).first()
    if not admin:
        print('No superuser found - cannot exercise the view. Aborting.')
        return 1

    span = Order.objects.filter(is_deleted=False).aggregate(first=Min('created_at'), last=Max('created_at'))
    if not span['first']:
        print('No orders in the local DB - nothing to verify against.')
        return 0
    today = timezone.localtime(timezone.now()).date()
    start = timezone.localtime(span['first']).date()
    end = min(timezone.localtime(span['last']).date(), today)
    PERIOD = dict(period='custom', custom_start=start.isoformat(), custom_end=end.isoformat())
    print(f'Acting as: {admin.username}   Period: custom ({start} .. {end})\n')

    print('1) orders_ids_only returns JSON ids matching the current filters')
    resp = call(admin, staff_filter='all', orders_ids_only='1', **PERIOD)
    check('HTTP 200', resp.status_code == 200, resp.status_code)
    check('content type is JSON', 'application/json' in resp['Content-Type'], resp['Content-Type'])
    data = json.loads(resp.content)
    ids = data.get('ids', [])

    from datetime import datetime, time as _time
    start_dt = timezone.make_aware(datetime.combine(start, _time.min))
    end_dt = timezone.make_aware(datetime.combine(end, _time.max))
    expected_qs = Order.objects.filter(created_at__gte=start_dt, created_at__lte=end_dt, is_deleted=False)
    check('id count matches filtered queryset', len(ids) == expected_qs.count(),
          f'{len(ids)} vs {expected_qs.count()}')
    check('ids are a subset of real order pks',
          set(ids) <= set(Order.objects.values_list('id', flat=True)))

    if len(ids) < 3:
        print('\nFewer than 3 orders in range - skipping selection-subset checks.')
        return 0 if not FAIL else 1

    print('\n2) orders_ids_only respects the status filter, same as the table')
    sample_order = Order.objects.filter(id__in=ids).select_related(None).first()
    a_status = sample_order.order_status or sample_order.status
    resp_status = call(admin, staff_filter='all', orders_ids_only='1', orders_status=a_status, **PERIOD)
    ids_status = json.loads(resp_status.content).get('ids', [])
    from django.db.models import Q
    expected_status_qs = expected_qs.filter(Q(status__iexact=a_status) | Q(order_status__iexact=a_status))
    check(f'status="{a_status}" id count matches', len(ids_status) == expected_status_qs.count(),
          f'{len(ids_status)} vs {expected_status_qs.count()}')

    print('\n3) order_ids export returns exactly the selected orders, in any order')
    picked = sorted(ids)[:3]  # simulate 3 checkboxes ticked across pages
    resp2 = call(admin, orders_export='xlsx', order_ids=','.join(str(i) for i in picked))
    check('HTTP 200', resp2.status_code == 200, resp2.status_code)
    wb = load_workbook(io.BytesIO(resp2.content))
    ws = wb['Staff Orders']
    exported_rows = ws.max_row - 1
    check('exactly 3 rows exported', exported_rows == 3, exported_rows)

    headers = [c.value for c in ws[1]]
    order_num_idx = headers.index('Order #')
    exported_numbers = {row[order_num_idx].value for row in ws.iter_rows(min_row=2)}
    expected_numbers = set(Order.objects.filter(id__in=picked).values_list('order_number', flat=True))
    check('exported order numbers match the picked ids exactly',
          exported_numbers == expected_numbers, f'{exported_numbers} vs {expected_numbers}')

    print('\n4) order_ids export ignores an unrelated status filter (explicit selection wins)')
    other_status_orders = expected_qs.exclude(id__in=picked).exclude(
        Q(status__iexact=a_status) | Q(order_status__iexact=a_status)
    )
    if other_status_orders.exists():
        mixed_ids = picked + [other_status_orders.first().id]
        resp3 = call(admin, orders_export='xlsx', order_ids=','.join(str(i) for i in mixed_ids),
                     orders_status=a_status, **PERIOD)
        ws3 = load_workbook(io.BytesIO(resp3.content))['Staff Orders']
        check('all requested ids present despite mismatched status filter',
              ws3.max_row - 1 == len(mixed_ids), f'{ws3.max_row - 1} vs {len(mixed_ids)}')

    print('\n5) Report Info sheet marks this as a manual selection, not a filtered export')
    info_ws = wb['Report Info']
    info = {r[0].value: r[1].value for r in info_ws.iter_rows(min_row=2)}
    check('Selection row present', 'Selection' in info, info)
    check('Selection note mentions the count', '3' in str(info.get('Selection', '')), info.get('Selection'))
    check('filename says "selection"',
          'staff_orders_selection' in resp2['Content-Disposition'], resp2['Content-Disposition'])

    print('\n6) Garbage / empty order_ids does not 500 - just yields an empty workbook')
    resp4 = call(admin, orders_export='xlsx', order_ids='not-a-number,,999999999')
    check('HTTP 200 on garbage ids', resp4.status_code == 200, resp4.status_code)
    ws4 = load_workbook(io.BytesIO(resp4.content))['Staff Orders']
    check('zero data rows for ids that do not exist', ws4.max_row == 1, ws4.max_row)

    print(f'\n{"=" * 60}\n{len(PASS)} passed, {len(FAIL)} failed')
    for f in FAIL:
        print('  FAILED: ' + f)
    return 1 if FAIL else 0


if __name__ == '__main__':
    sys.exit(main())
