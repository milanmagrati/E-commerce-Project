"""Verify the Bulk Logs *list* export (batches / order rows / both, xlsx + csv).

Run:  python test_bulk_log_list_export.py

Companion to test_bulk_log_export.py, which covers the single-batch export.
This one drives the list page's own export URL through Django's test client
against the live DB, so it proves the shared filter helpers, the scope and
include switches, the download headers and the permission decorator are wired
up - not just the row builders.

Most of the interesting shapes (two providers in one file, mixed statuses,
Nepali text, a batch with no rows) do not exist in a developer database, so the
second half builds them inside a transaction that is rolled back.
"""
import os
import csv
import io

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'myproject.settings')
django.setup()

from django.conf import settings  # noqa: E402

# The test client speaks to host "testserver", which the project's real
# ALLOWED_HOSTS has no reason to contain.
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS = list(settings.ALLOWED_HOSTS) + ['testserver']

from django.db import transaction  # noqa: E402
from django.test import Client  # noqa: E402
from django.test.utils import setup_test_environment  # noqa: E402
from django.urls import reverse  # noqa: E402

# response.context is only captured when the instrumented template renderer is
# installed; a standalone script has to ask for it explicitly.
try:
    setup_test_environment()
except RuntimeError:
    pass
from openpyxl import load_workbook  # noqa: E402

from accounts.models import CustomUser  # noqa: E402
from ncm.models import NCMBulkLog, NCMBulkLogOrder  # noqa: E402
from pick_and_drop.models import PNDBulkLog, PNDBulkLogOrder  # noqa: E402
from dashboard.views import (  # noqa: E402
    BULK_LOG_LIST_NO_FILTERS,
    _bulk_log_export_columns,
    _bulk_log_list_batch_columns,
    _bulk_log_list_batch_rows,
    _bulk_log_list_batches,
    _bulk_log_list_order_columns,
    _bulk_log_list_order_counts,
    _bulk_log_list_order_rows,
)

failures = []

EXPORT_URL = reverse('logistics_bulk_logs_export')
LIST_URL = reverse('logistics_bulk_logs_list')


def check(label, condition, detail=''):
    status = 'PASS' if condition else 'FAIL'
    print('  [{}] {}{}'.format(status, label, ' - ' + detail if detail else ''))
    if not condition:
        failures.append(label)


def sheet_rows(sheet):
    return [[c.value for c in row] for row in sheet.iter_rows(min_row=2)]


def parse_csv(response):
    return list(csv.reader(io.StringIO(response.content.decode('utf-8-sig'))))


# -- 1. Columns: the flattened sheet is the batch sheet plus its provenance --
print('\n=== Column shapes ===')
base_columns = _bulk_log_export_columns('Tracking ID')
order_columns = _bulk_log_list_order_columns()
check('order columns are the batch export columns plus Batch Number / Provider',
      order_columns == base_columns[:1] + ['Batch Number', 'Provider'] + base_columns[1:],
      str(order_columns))
check('order columns still lead with the running number', order_columns[0] == 'S.N.')
check('batch columns cover the list table and its stored counters',
      set(['Batch Number', 'Provider', 'Status', 'Total Orders', 'Success', 'Failed',
           'Skipped', 'Branch', 'Sent By', 'Sent At']).issubset(_bulk_log_list_batch_columns()),
      str(_bulk_log_list_batch_columns()))


# -- 2. The shared lookup the page and the export both read -----------------
print('\n=== Shared batch lookup ===')
batches, branches, ncm_all, pnd_all = _bulk_log_list_batches('all', dict(BULK_LOG_LIST_NO_FILTERS))
check('every batch is tagged with its provider',
      all(b.log_provider in ('ncm', 'pnd') for b in batches))
check('every batch carries a branch for display',
      all(hasattr(b, 'branch_display') for b in batches))
check('batches come back newest first',
      all(batches[i].created_at >= batches[i + 1].created_at for i in range(len(batches) - 1)))
check('the unfiltered querysets match the provider tables',
      ncm_all.count() == NCMBulkLog.objects.filter(is_deleted=False).count()
      and pnd_all.count() == PNDBulkLog.objects.filter(is_deleted=False).count())
check('branch choices are unique and sorted', branches == sorted(set(branches)))

ncm_only, _b, _n, _p = _bulk_log_list_batches('ncm', dict(BULK_LOG_LIST_NO_FILTERS))
check('provider=ncm returns only NCM batches',
      all(b.log_provider == 'ncm' for b in ncm_only))
check('provider=ncm matches the NCM row count',
      len(ncm_only) == NCMBulkLog.objects.filter(is_deleted=False).count(),
      '{} vs {}'.format(len(ncm_only), NCMBulkLog.objects.filter(is_deleted=False).count()))

nonsense = dict(BULK_LOG_LIST_NO_FILTERS, search='zzz-no-such-batch-zzz')
check('a filter that matches nothing returns nothing',
      _bulk_log_list_batches('all', nonsense)[0] == [])

counts = _bulk_log_list_order_counts(batches)
recorded = (NCMBulkLogOrder.objects.filter(batch__is_deleted=False).count()
            + PNDBulkLogOrder.objects.filter(batch__is_deleted=False).count())
check('order counts are read from the rows, not the batch counters',
      counts['all'] == recorded, '{} vs {}'.format(counts['all'], recorded))
check('the three statuses add up to the total',
      counts['success'] + counts['failed'] + counts['skipped'] == counts['all'])
check('the batch count comes along for the dialog', counts['batches'] == len(batches))

rows, truncated = _bulk_log_list_order_rows(batches, [])
check('order rows cover every recorded row', len(rows) == recorded,
      '{} vs {}'.format(len(rows), recorded))
check('nothing is truncated at this size', truncated is False)
check('every order row has one cell per column',
      all(len(r) == len(order_columns) for r in rows))
check('order rows are numbered once across the whole file',
      [r[0] for r in rows] == list(range(1, len(rows) + 1)))
batch_numbers = {b.batch_number for b in batches}
check('every order row names a batch from the list',
      all(r[1] in batch_numbers for r in rows))

batch_rows = _bulk_log_list_batch_rows(batches)
check('one batch row per batch', len(batch_rows) == len(batches))
check('every batch row has one cell per column',
      all(len(r) == len(_bulk_log_list_batch_columns()) for r in batch_rows))


# -- 3. The view itself, through the URL, as a logged-in admin --------------
admin = (CustomUser.objects.filter(is_superuser=True).first()
         or CustomUser.objects.filter(role='administrator').first())

print('\n=== HTTP responses ===')
if admin is None:
    print('  [SKIP] no administrator user in this database')
    client = None
else:
    client = Client()
    client.force_login(admin)

    xlsx = client.get(EXPORT_URL, {'include': 'both', 'format': 'xlsx'})
    check('xlsx returns 200', xlsx.status_code == 200, str(xlsx.status_code))
    check('xlsx has the spreadsheet content type', 'spreadsheetml' in xlsx['Content-Type'])
    check('xlsx is an attachment with a dated name',
          xlsx['Content-Disposition'].startswith('attachment')
          and 'bulk-logs-all-both-' in xlsx['Content-Disposition'],
          xlsx.get('Content-Disposition', ''))

    wb = load_workbook(io.BytesIO(xlsx.content))
    check('both-sheet workbook holds Summary, Batches and Orders',
          wb.sheetnames == ['Summary', 'Batches', 'Orders'], str(wb.sheetnames))
    check('batch sheet header matches the column list',
          [c.value for c in wb['Batches'][1]] == _bulk_log_list_batch_columns())
    check('order sheet header matches the column list',
          [c.value for c in wb['Orders'][1]] == order_columns)
    check('batch sheet has one row per batch',
          wb['Batches'].max_row - 1 == len(batches),
          '{} vs {}'.format(wb['Batches'].max_row - 1, len(batches)))
    check('order sheet has one row per order row',
          wb['Orders'].max_row - 1 == len(rows),
          '{} vs {}'.format(wb['Orders'].max_row - 1, len(rows)))
    check('both sheets freeze their header',
          wb['Batches'].freeze_panes == 'A2' and wb['Orders'].freeze_panes == 'A2')

    summary = {row[0]: row[1] for row in wb['Summary'].iter_rows(values_only=True)}
    check('summary states the scope, the filters and the counts',
          summary.get('Scope') == 'Current view'
          and summary.get('Filters') == 'None'
          and summary.get('Batches') == len(batches),
          str(summary))
    check('summary names who generated the file',
          summary.get('Generated By') == admin.get_username(), str(summary.get('Generated By')))

    batches_only = load_workbook(io.BytesIO(
        client.get(EXPORT_URL, {'include': 'batches'}).content))
    check('a batches-only workbook has no Orders sheet',
          batches_only.sheetnames == ['Summary', 'Batches'], str(batches_only.sheetnames))
    orders_only = load_workbook(io.BytesIO(
        client.get(EXPORT_URL, {'include': 'orders'}).content))
    check('an orders-only workbook has no Batches sheet',
          orders_only.sheetnames == ['Summary', 'Orders'], str(orders_only.sheetnames))

    csv_orders = client.get(EXPORT_URL, {'include': 'orders', 'format': 'csv'})
    check('csv returns 200', csv_orders.status_code == 200, str(csv_orders.status_code))
    check('csv is BOM-prefixed for Excel', csv_orders.content.startswith(b'\xef\xbb\xbf'))
    parsed = parse_csv(csv_orders)
    check('csv header matches the order columns', parsed[0] == order_columns)
    check('csv body matches the order rows', len(parsed) - 1 == len(rows),
          '{} vs {}'.format(len(parsed) - 1, len(rows)))

    csv_both = client.get(EXPORT_URL, {'include': 'both', 'format': 'csv'})
    check('a csv asking for both tables falls back to the batch list',
          parse_csv(csv_both)[0] == _bulk_log_list_batch_columns()
          and 'bulk-logs-all-batches-' in csv_both['Content-Disposition'],
          csv_both.get('Content-Disposition', ''))

    junk = client.get(EXPORT_URL, {'include': 'nonsense', 'format': 'nonsense',
                                   'scope': 'nonsense', 'provider': 'nonsense'})
    check('nonsense parameters fall back instead of erroring',
          junk.status_code == 200 and 'spreadsheetml' in junk['Content-Type']
          and 'bulk-logs-all-batches-' in junk['Content-Disposition'],
          junk.get('Content-Disposition', ''))

    empty_selection = client.get(EXPORT_URL, {'scope': 'selected', 'ids': ''})
    check('selecting nothing exports an empty batch list, not an error',
          empty_selection.status_code == 200
          and load_workbook(io.BytesIO(empty_selection.content))['Batches'].max_row == 1,
          str(empty_selection.status_code))

    anon = Client().get(EXPORT_URL)
    check('anonymous request is not served the file',
          anon.status_code in (302, 401, 403), str(anon.status_code))

    print('\n=== List page wiring ===')
    page = client.get(LIST_URL)
    check('list page renders', page.status_code == 200, str(page.status_code))
    body = page.content.decode('utf-8', 'ignore')
    check('list page carries the export dialog', 'id="exportLogsModal"' in body)
    check('list page posts the dialog at the export URL', EXPORT_URL in body)
    check('page counts match the exportable rows',
          page.context['export_counts'] == counts,
          '{} vs {}'.format(page.context.get('export_counts'), counts))
    check('page reports how many batches the current view holds',
          page.context['filtered_batches'] == len(batches),
          '{} vs {}'.format(page.context.get('filtered_batches'), len(batches)))
    check('an unfiltered page says so', page.context['has_filters'] is False)
    filtered_page = client.get(LIST_URL, {'search': 'zzz-no-such-batch-zzz'})
    check('a filtered page says so', filtered_page.context['has_filters'] is True)
    check('a filter that matches nothing leaves no batches to export',
          filtered_page.context['filtered_batches'] == 0)


class _Rollback(Exception):
    """Thrown to unwind the synthetic batches - nothing below is persisted."""


def run_synthetic(admin):
    """Two batches, one per provider, with the mixed statuses and Nepali text a
    developer database has no reason to contain.

    Written inside a transaction that is rolled back, so the database is left
    exactly as it was found.
    """
    print('\n=== Synthetic batches ===')
    if admin is None:
        print('  [SKIP] no administrator user in this database')
        return

    client = Client()
    client.force_login(admin)
    marker = 'ZZTEST'

    try:
        with transaction.atomic():
            ncm_log = NCMBulkLog.objects.create(
                batch_number='NCM-' + marker, total_orders=3, success_count=1,
                failed_count=1, skipped_count=1, status='partial', created_by=admin,
                from_branch='TINKUNE', delivery_type='Door2Door')
            pnd_log = PNDBulkLog.objects.create(
                batch_number='PND-' + marker, total_orders=1, success_count=0,
                failed_count=1, status='failed', created_by=admin,
                destination_branch='KATHMANDU VALLEY')

            for i, (status, tracking) in enumerate(
                    [('success', 25137563), ('failed', None), ('skipped', None)], start=1):
                NCMBulkLogOrder.objects.create(
                    batch=ncm_log, order_number='{}-{}'.format(marker, i),
                    customer_name='सबिता श्रेष्ठ', customer_phone='980000000{}'.format(i),
                    shipping_address='काठमाडौं, नयाँ बसपार्क', cod_amount=1500,
                    destination_branch='TINKUNE', status=status,
                    message='Rejected\x01by API' if status == 'failed' else 'ok',
                    ncm_order_id=tracking)
            PNDBulkLogOrder.objects.create(
                batch=pnd_log, order_number=marker + '-P1', customer_name='कबिता',
                customer_phone='9800000009', shipping_address='ललितपुर',
                cod_amount=900, destination_branch='KATHMANDU VALLEY',
                status='failed', message='No response', pnd_order_id=None)

            search = {'search': marker}
            synthetic, _b, _n, _p = _bulk_log_list_batches('all', dict(
                BULK_LOG_LIST_NO_FILTERS, **search))
            check('both synthetic batches land in one filtered view',
                  {b.batch_number for b in synthetic} == {'NCM-' + marker, 'PND-' + marker},
                  str([b.batch_number for b in synthetic]))

            wb = load_workbook(io.BytesIO(client.get(
                EXPORT_URL, dict(search, include='both')).content))
            batch_sheet = sheet_rows(wb['Batches'])
            order_sheet = sheet_rows(wb['Orders'])
            batch_cols = _bulk_log_list_batch_columns()

            check('the filtered export carries exactly the two batches',
                  len(batch_sheet) == 2, str(len(batch_sheet)))
            check('the export names both providers in words',
                  {r[batch_cols.index('Provider')] for r in batch_sheet}
                  == {'Nepal Can Move', 'Pick & Drop'},
                  str([r[batch_cols.index('Provider')] for r in batch_sheet]))
            check('a Pick & Drop batch exports a blank delivery type, not "None"',
                  [r[batch_cols.index('Delivery Type')] for r in batch_sheet
                   if r[batch_cols.index('Provider')] == 'Pick & Drop'] == [None],
                  str(batch_sheet))
            check('both providers\' order rows share one sheet',
                  len(order_sheet) == 4, str(len(order_sheet)))
            check('order rows are numbered across providers',
                  [r[0] for r in order_sheet] == [1, 2, 3, 4],
                  str([r[0] for r in order_sheet]))
            check('each order row names the batch it came from',
                  {r[1] for r in order_sheet} == {'NCM-' + marker, 'PND-' + marker},
                  str([r[1] for r in order_sheet]))
            check('a missing tracking id exports blank, not "None"',
                  order_sheet[-1][order_columns.index('Tracking ID')] in (None, ''),
                  repr(order_sheet[-1][order_columns.index('Tracking ID')]))
            check('the control character never reaches the workbook',
                  all('\x01' not in (r[order_columns.index('Message')] or '')
                      for r in order_sheet))

            status_col = order_columns.index('Status')
            fills = {r: wb['Orders'].cell(row=r, column=1).fill.start_color.rgb
                     for r in range(2, wb['Orders'].max_row + 1)}
            statuses = {wb['Orders'].cell(row=r, column=status_col + 1).value: fills[r]
                        for r in fills}
            check('failed rows are tinted red',
                  str(statuses.get('Failed')).endswith('FFE4E6'), str(statuses.get('Failed')))
            check('skipped rows are tinted amber',
                  str(statuses.get('Skipped')).endswith('FEF3C7'), str(statuses.get('Skipped')))
            check('success rows are left untinted',
                  not str(statuses.get('Success')).endswith(('FFE4E6', 'FEF3C7')),
                  str(statuses.get('Success')))

            failed_only = load_workbook(io.BytesIO(client.get(
                EXPORT_URL, dict(search, include='orders', order_status='failed')).content))
            failed_rows = sheet_rows(failed_only['Orders'])
            check('a failed-only export carries just the failed rows',
                  len(failed_rows) == 2 and all(r[status_col] == 'Failed' for r in failed_rows),
                  str([r[status_col] for r in failed_rows]))
            failed_summary = {row[0]: row[1] for row
                              in failed_only['Summary'].iter_rows(values_only=True)}
            check('the summary records which statuses were asked for',
                  failed_summary.get('Order Statuses') == 'Failed',
                  str(failed_summary.get('Order Statuses')))
            check('the summary records the filter that was applied',
                  marker in str(failed_summary.get('Filters')),
                  str(failed_summary.get('Filters')))

            two_status = load_workbook(io.BytesIO(client.get(
                EXPORT_URL, dict(search, include='orders',
                                 order_status=['failed', 'skipped'])).content))
            check('two ticked statuses come out together',
                  {r[status_col] for r in sheet_rows(two_status['Orders'])}
                  == {'Failed', 'Skipped'})

            ncm_provider = load_workbook(io.BytesIO(client.get(
                EXPORT_URL, dict(search, provider='ncm', include='batches')).content))
            check('provider=ncm drops the Pick & Drop batch',
                  [r[batch_cols.index('Batch Number')]
                   for r in sheet_rows(ncm_provider['Batches'])] == ['NCM-' + marker])

            selected = load_workbook(io.BytesIO(client.get(EXPORT_URL, {
                'scope': 'selected', 'ids': 'pnd:{}'.format(pnd_log.id),
                'include': 'both'}).content))
            check('scope=selected exports only the ticked batch',
                  [r[batch_cols.index('Batch Number')]
                   for r in sheet_rows(selected['Batches'])] == ['PND-' + marker],
                  str(sheet_rows(selected['Batches'])))
            check('scope=selected pulls that batch\'s order rows with it',
                  len(sheet_rows(selected['Orders'])) == 1)
            selected_summary = {row[0]: row[1] for row
                                in selected['Summary'].iter_rows(values_only=True)}
            check('the summary says the filters were not used',
                  selected_summary.get('Scope') == 'Selected batches'
                  and selected_summary.get('Filters') == 'None',
                  str(selected_summary))

            all_scope = load_workbook(io.BytesIO(client.get(
                EXPORT_URL, dict(search, scope='all', include='batches')).content))
            check('scope=all ignores the filter and takes every batch',
                  all_scope['Batches'].max_row - 1
                  == NCMBulkLog.objects.filter(is_deleted=False).count()
                  + PNDBulkLog.objects.filter(is_deleted=False).count(),
                  str(all_scope['Batches'].max_row - 1))

            nepali_csv = client.get(EXPORT_URL, dict(search, include='orders', format='csv'))
            check('Nepali text survives the csv round-trip',
                  'सबिता श्रेष्ठ' in nepali_csv.content.decode('utf-8-sig'))

            raise _Rollback
    except _Rollback:
        pass

    check('synthetic batches left no trace',
          not NCMBulkLog.objects.filter(batch_number='NCM-' + marker).exists()
          and not PNDBulkLog.objects.filter(batch_number='PND-' + marker).exists())


run_synthetic(admin)

print('\n' + ('ALL CHECKS PASSED' if not failures else 'FAILED: ' + ', '.join(failures)))
raise SystemExit(1 if failures else 0)
