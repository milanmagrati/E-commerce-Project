"""Verify the single-batch dispatch export (Excel + CSV, all three scopes).

Run:  python test_dispatch_export.py

Exercises the real view through Django's test client against the live DB, so
it also proves the URL, the permission decorator and the download headers are
wired up — not just the row-building helpers.
"""
import os
import csv
import io

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'myproject.settings')
django.setup()

from django.conf import settings  # noqa: E402

# The test client speaks to host "testserver", which the project's real
# ALLOWED_HOSTS has no reason to contain.
if 'testserver' not in settings.ALLOWED_HOSTS:
    settings.ALLOWED_HOSTS = list(settings.ALLOWED_HOSTS) + ['testserver']

from django.test import Client  # noqa: E402
from django.urls import reverse  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

from accounts.models import CustomUser  # noqa: E402
from dashboard.models import Dispatch  # noqa: E402
from dashboard.views import (  # noqa: E402
    DISPATCH_EXPORT_COLUMNS,
    _dispatch_export_rows,
)

failures = []


def check(label, condition, detail=''):
    status = 'PASS' if condition else 'FAIL'
    print('  [{}] {}{}'.format(status, label, ' - ' + detail if detail else ''))
    if not condition:
        failures.append(label)


# ── Pick the most useful batch available: prefer one that has problem rows,
#    since that is the scope with the most branching in the row builder. ──────
dispatch = None
for candidate in Dispatch.objects.prefetch_related('items__order').order_by('-created_at')[:50]:
    if candidate.get_issue_count():
        dispatch = candidate
        break
if dispatch is None:
    dispatch = Dispatch.objects.prefetch_related('items__order').order_by('-created_at').first()

if dispatch is None:
    print('No Dispatch rows in this database - nothing to export. Skipping.')
    raise SystemExit(0)

print('\nBatch under test: {} ({} scans: {} ok / {} failed / {} not found)'.format(
    dispatch.batch_number,
    dispatch.get_item_count(),
    dispatch.get_success_count(),
    dispatch.get_failed_count(),
    dispatch.get_not_found_count(),
))

# ── 1. Row builder: scopes partition the scans, columns line up ──────────────
print('\n1. Row builder')
rows_all = _dispatch_export_rows(dispatch, 'all')
rows_ok = _dispatch_export_rows(dispatch, 'success')
rows_bad = _dispatch_export_rows(dispatch, 'problems')

check('all scope covers every recorded scan',
      len(rows_all) == dispatch.get_item_count(),
      '{} rows vs {} items'.format(len(rows_all), dispatch.get_item_count()))
check('success scope matches success count',
      len(rows_ok) == dispatch.get_success_count())
check('problems scope matches issue count',
      len(rows_bad) == dispatch.get_issue_count())
check('success + problems = all', len(rows_ok) + len(rows_bad) == len(rows_all))
check('every row has one cell per column',
      all(len(r) == len(DISPATCH_EXPORT_COLUMNS) for r in rows_all))
check('row numbering restarts at 1 per scope',
      not rows_bad or rows_bad[0][0] == 1)

outcome_idx = DISPATCH_EXPORT_COLUMNS.index('Outcome')
reason_idx = DISPATCH_EXPORT_COLUMNS.index('Reason')
check('no success row leaks into the problems scope',
      all(r[outcome_idx] != 'Success' for r in rows_bad))
check('every problem row carries a reason',
      all(str(r[reason_idx]).strip() for r in rows_bad))

# ── 2. The view itself, through the URL, as a logged-in admin ────────────────
print('\n2. HTTP responses')
admin = CustomUser.objects.filter(is_superuser=True).first() or \
    CustomUser.objects.filter(role='administrator').first()
if admin is None:
    print('  [SKIP] no administrator user in this database')
else:
    client = Client()
    client.force_login(admin)
    url = reverse('dispatch_export', args=[dispatch.pk])

    xlsx = client.get(url, {'format': 'xlsx', 'scope': 'all'})
    check('xlsx returns 200', xlsx.status_code == 200, str(xlsx.status_code))
    check('xlsx is an attachment named after the batch',
          dispatch.batch_number in xlsx['Content-Disposition']
          and xlsx['Content-Disposition'].startswith('attachment'),
          xlsx.get('Content-Disposition', ''))
    check('xlsx has the spreadsheet content type',
          'spreadsheetml' in xlsx['Content-Type'])

    wb = load_workbook(io.BytesIO(xlsx.content))
    check('workbook has Summary + Scanned Orders sheets',
          wb.sheetnames == ['Summary', 'Scanned Orders'], str(wb.sheetnames))
    sheet = wb['Scanned Orders']
    check('sheet header matches the shared column list',
          [c.value for c in sheet[1]] == DISPATCH_EXPORT_COLUMNS)
    check('sheet has one data row per scan',
          sheet.max_row - 1 == len(rows_all),
          '{} data rows vs {} scans'.format(sheet.max_row - 1, len(rows_all)))
    check('header row is frozen', sheet.freeze_panes == 'A2')

    csv_resp = client.get(url, {'format': 'csv', 'scope': 'problems'})
    check('csv returns 200', csv_resp.status_code == 200, str(csv_resp.status_code))
    check('csv filename reflects the scope',
          'problems.csv' in csv_resp['Content-Disposition'],
          csv_resp.get('Content-Disposition', ''))
    body = csv_resp.content.decode('utf-8-sig')
    check('csv is BOM-prefixed for Excel', csv_resp.content.startswith(b'\xef\xbb\xbf'))
    parsed = list(csv.reader(io.StringIO(body)))
    check('csv header matches the shared column list',
          parsed[0] == DISPATCH_EXPORT_COLUMNS)
    check('csv body matches the problems scope',
          len(parsed) - 1 == len(rows_bad),
          '{} rows vs {} problems'.format(len(parsed) - 1, len(rows_bad)))

    bad_scope = client.get(url, {'format': 'xlsx', 'scope': 'nonsense'})
    check('unknown scope falls back to all instead of erroring',
          bad_scope.status_code == 200
          and 'all-scans' in bad_scope['Content-Disposition'],
          bad_scope.get('Content-Disposition', ''))

    default_fmt = client.get(url)
    check('no format defaults to xlsx',
          default_fmt.status_code == 200
          and 'spreadsheetml' in default_fmt['Content-Type'])

    anon = Client().get(url)
    check('anonymous request is not served the file',
          anon.status_code in (302, 401, 403), str(anon.status_code))

print('\n' + ('ALL CHECKS PASSED' if not failures else 'FAILED: ' + ', '.join(failures)))
raise SystemExit(1 if failures else 0)
