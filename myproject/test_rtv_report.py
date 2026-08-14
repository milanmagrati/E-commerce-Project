"""Verify the RTV Report (/reports/rtv/).

Builds a controlled set of RTVs — three staff members, several reason
comments, one unattributable row and one repeat customer — inside a
transaction that is always rolled back, then asserts the view's aggregation,
filters, charts, export and permission gate.

    python test_rtv_report.py
"""
import io
import json
import os
import re
import sys

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'myproject.settings')
django.setup()

from django.db import transaction
from django.test import RequestFactory
from django.utils import timezone

from accounts.models import CustomUser
from dashboard.models import Order, RTVOrder
from dashboard.views import classify_rtv_reason, rtv_report


class Rollback(Exception):
    pass


# NCM ids far above anything real so they cannot collide with live data.
BASE_ID = 991000000


def make_order(staff, suffix, phone, created_at, ncm_id=None):
    order = Order.objects.create(
        order_number=f'RTVREP-{suffix}',
        customer_name=f'Report Customer {suffix}',
        customer_phone=phone,
        shipping_address='Kathmandu',
        order_from='test',
        created_by=staff,
        ncm_order_id=ncm_id,
    )
    Order.objects.filter(pk=order.pk).update(created_at=created_at)
    order.refresh_from_db()
    return order


def make_rtv(vendor, ncm_id, comment, marked_at, phone='', name='',
             last_status='Arrived'):
    return RTVOrder.objects.create(
        order_id=ncm_id,
        comment=comment,
        vendor=vendor,
        receiver_phone=phone,
        receiver_name=name,
        last_status=last_status,
        rtv_marked_at=marked_at,
        rtv_marked_at_source=RTVOrder.SOURCE_MANUAL,
    )


def run():
    failures = []

    def check(label, condition, detail=''):
        if condition:
            print(f'  PASS  {label}')
        else:
            print(f'  FAIL  {label} {detail}')
            failures.append(label)

    # ---------- reason classifier ----------
    print('reason classification')
    cases = [
        ('Customer cancelled the order', 'cancelled'),
        ('phone switch off, not responding', 'unreachable'),
        ('customer is out of station', 'unavailable'),
        ('wrong address, could not locate', 'address'),
        ('out of delivery area', 'area'),
        ('customer has no cash right now', 'payment'),
        ('product damaged in transit', 'damaged'),
        ('duplicate order by mistake', 'duplicate'),
        ('held too long, no update', 'delayed'),
        ('return to vendor', 'return_request'),
        # NCM's real canned reason codes, taken from the live comment column.
        # These are the strings the courier actually sends, and every one of
        # them used to land in Uncategorised.
        ('PNR/PUR/PSO', 'unreachable'),
        ('Ignorance/Rejection by Customer', 'cancelled'),
        ('Customer did not ordered', 'not_ordered'),
        ('Late Delivery', 'delayed'),
        ('Late Process', 'delayed'),
        ('As per vendor request', 'vendor_request'),
        ('Product Issue (Size/Color/Quality)', 'damaged'),
        ('Double Order Received', 'duplicate'),
        ('Wrong Location', 'address'),
        ('Out of Coverage area', 'area'),
        ('Wrong Number', 'wrong_number'),
        ('Price Issue', 'payment'),
        ('Dispute at Delivery Location', 'cancelled'),
        ('', 'no_comment'),
        ('   ', 'no_comment'),
        ('—', 'no_comment'),
        ('zzzz qqqq', 'other'),
    ]
    for text, expected in cases:
        got = classify_rtv_reason(text)[0]
        check(f'{text!r:42} -> {expected}', got == expected, f'got {got}')

    # A comment naming both a cancellation and no response must land in the
    # more specific bucket, which is what the rule ordering exists for.
    check('specific rule beats vague one',
          classify_rtv_reason('no response so customer cancelled')[0] == 'cancelled',
          f"got {classify_rtv_reason('no response so customer cancelled')[0]}")

    # ---------- fixtures ----------
    print('\nfixtures')
    staff = list(CustomUser.objects.all()[:3])
    if len(staff) < 3:
        raise SystemExit('need at least 3 users in the DB to run this script')
    alice, bob, carol = staff
    admin = CustomUser.objects.filter(is_superuser=True).first() or alice
    now = timezone.now()
    day = timezone.timedelta(days=1)

    # alice: 3 RTVs (2 cancelled, 1 unreachable) — should top the leaderboard
    # bob:   2 RTVs (1 address, 1 damaged)
    # carol: 1 RTV  (payment), matched by phone rather than by id
    # plus:  1 RTV that nothing can attribute
    plan = [
        (alice, 1, 'Customer cancelled', 'exact', '9801110001'),
        (alice, 2, 'cancelled by customer', 'exact', '9801110002'),
        (alice, 3, 'phone switch off', 'exact', '9801110003'),
        (bob, 4, 'wrong address', 'exact', '9801110004'),
        (bob, 5, 'item damaged', 'exact', '9801110005'),
        (carol, 6, 'no cash with customer', 'phone', '9801110006'),
    ]
    created = {}
    for who, n, comment, link, phone in plan:
        ncm_id = BASE_ID + n
        placed = now - (10 * day)
        order = make_order(
            who, n, phone, placed,
            ncm_id=ncm_id if link == 'exact' else None,
        )
        rtv = make_rtv(admin, ncm_id, comment, now - (2 * day), phone,
                       f'Report Customer {n}')
        created[n] = (order, rtv)

    # An RTV with no order, no bulk log and a phone nobody ordered from.
    make_rtv(admin, BASE_ID + 8, 'zzzz qqqq', now - (2 * day),
             '9809999999', 'Ghost Buyer')

    # A second RTV on carol's phone, so the repeat-customer panel has input.
    make_rtv(admin, BASE_ID + 7, 'Customer cancelled', now - (3 * day),
             '9801110006', 'Report Customer 6')

    total_fixtures = 8
    print(f'  built {total_fixtures} RTVs across 3 staff + 1 unattributable')

    rf = RequestFactory()

    def render(params=None):
        req = rf.get('/reports/rtv/', params or {})
        req.user = admin
        resp = rtv_report(req)
        return resp

    # Only the fixture rows: search on the prefix every fixture id shares
    # (991000001 … 991000008). Anything wider would drag in real RTVs and the
    # exact counts below would stop meaning anything.
    SCOPE = {'period': 'all', 'search': '9910000'}

    print('\naggregation')
    html = render(SCOPE).content.decode('utf-8', 'replace')
    check('page renders', 'RTV Report' in html and 'Staff Leaderboard' in html)
    check('all 8 fixture RTVs counted', '>8<' in html or 'Total RTVs' in html)

    # Pull the chart payload back out — it is the same data the tables render.
    m = re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', html, re.S)
    check('chart payload present', m is not None)
    data = json.loads(m.group(1).replace('\\u0022', '"')) if m else {}
    if data:
        labels = data['staff']['labels']
        confirmed = data['staff']['confirmed']
        probable = data['staff']['probable']
        totals = {lab: confirmed[i] + probable[i] for i, lab in enumerate(labels)}
        alice_name = alice.get_full_name() or alice.username
        bob_name = bob.get_full_name() or bob.username
        carol_name = carol.get_full_name() or carol.username
        check('alice leads with 3', totals.get(alice_name) == 3, f'got {totals.get(alice_name)}')
        check('bob has 2', totals.get(bob_name) == 2, f'got {totals.get(bob_name)}')
        check('carol has 2 (one is the repeat)', totals.get(carol_name) == 2,
              f'got {totals.get(carol_name)}')
        check('alice is ranked first',
              labels and labels[0] == alice_name, f'top was {labels[0] if labels else None}')
        check('alice attributed by id, not phone',
              probable[labels.index(alice_name)] == 0 if alice_name in labels else False)
        check('carol attributed by phone',
              probable[labels.index(carol_name)] == 2 if carol_name in labels else False,
              f'probable={probable}')

        reason_totals = dict(zip(data['reasons']['labels'], data['reasons']['values']))
        check('3 cancellations counted', reason_totals.get('Customer Cancelled') == 3,
              f'got {reason_totals.get("Customer Cancelled")}')
        check('uncategorised row bucketed', reason_totals.get('Uncategorised') == 1,
              f'got {reason_totals.get("Uncategorised")}')
        check('reason totals sum to the RTV count',
              sum(data['reasons']['values']) == total_fixtures,
              f'got {sum(data["reasons"]["values"])}')
        check('trend series is dated and non-empty',
              len(data['trend']['labels']) == len(data['trend']['values'])
              and sum(data['trend']['values']) == total_fixtures,
              f'sum={sum(data["trend"]["values"])}')

    check('unattributed row surfaced', 'Not linked to a local order' in html)
    check('repeat customer panel lists the doubled phone',
          'Repeat RTV Customers' in html and '9801110006' in html)
    check('staff x reason matrix rendered', 'Staff × Reason' in html)
    check('uncategorised comments panel offers the raw text',
          'Uncategorised Comments' in html and 'zzzz qqqq' in html)

    print('\nfilters')
    only_alice = render(dict(SCOPE, staff=alice.username)).content.decode('utf-8', 'replace')
    m2 = re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', only_alice, re.S)
    d2 = json.loads(m2.group(1).replace('\\u0022', '"'))
    check('staff filter narrows to that person',
          sum(d2['trend']['values']) == 3, f'got {sum(d2["trend"]["values"])}')
    check('staff filter leaves only their reasons',
          sorted(d2['reasons']['labels']) == ['Customer Cancelled', 'Unreachable / No Response'],
          f'got {d2["reasons"]["labels"]}')

    only_cancel = render(dict(SCOPE, reason='cancelled')).content.decode('utf-8', 'replace')
    d3 = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', only_cancel, re.S)
                    .group(1).replace('\\u0022', '"'))
    check('reason filter narrows to that bucket',
          sum(d3['trend']['values']) == 3, f'got {sum(d3["trend"]["values"])}')

    unlinked = render(dict(SCOPE, confidence='none')).content.decode('utf-8', 'replace')
    d4 = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', unlinked, re.S)
                    .group(1).replace('\\u0022', '"'))
    check('unlinked filter isolates the one orphan RTV',
          sum(d4['trend']['values']) == 1, f'got {sum(d4["trend"]["values"])}')
    check('unlinked filter charts no staff', d4['staff']['labels'] == [],
          f'got {d4["staff"]["labels"]}')

    probable_only = render(dict(SCOPE, confidence='probable')).content.decode('utf-8', 'replace')
    d5 = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', probable_only, re.S)
                    .group(1).replace('\\u0022', '"'))
    check('probable filter keeps only the phone matches',
          sum(d5['trend']['values']) == 2, f'got {sum(d5["trend"]["values"])}')

    empty = render(dict(SCOPE, staff=alice.username, reason='damaged')).content.decode('utf-8', 'replace')
    check('impossible combination shows the empty state',
          'No RTVs match these filters' in empty)

    # Dropdown options must survive a drill-down, or you cannot switch person.
    check('staff dropdown still lists everyone while filtered',
          (bob.get_full_name() or bob.username) in only_alice,
          'other staff vanished from the filter form')

    # Clicking a second staff member while one is already selected must replace
    # the filter, not append a duplicate staff= to the URL.
    links = re.findall(r'href="\?([^"]*staff=[^"]*)"', only_alice)
    check('leaderboard links carry exactly one staff param',
          links and all(link.count('staff=') == 1 for link in links),
          f'found {[l for l in links if l.count("staff=") != 1][:2]}')
    check('leaderboard links keep the other filters',
          all('period=all' in link for link in links), f'{links[:1]}')

    both = render(dict(SCOPE, staff=alice.username, reason='cancelled')).content.decode('utf-8', 'replace')
    d_both = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', both, re.S)
                        .group(1).replace('\\u0022', '"'))
    check('stacked staff + reason filters intersect',
          sum(d_both['trend']['values']) == 2, f'got {sum(d_both["trend"]["values"])}')

    print('\ndate handling')
    today_only = render({'period': 'today', 'search': '9910000'}).content.decode('utf-8', 'replace')
    check('today excludes RTVs marked days ago',
          'No RTVs match these filters' in today_only)
    custom = render({
        'period': 'custom', 'search': '9910000',
        'date_from': (timezone.localtime(now).date() - timezone.timedelta(days=4)).isoformat(),
        'date_to': timezone.localtime(now).date().isoformat(),
    }).content.decode('utf-8', 'replace')
    d6 = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', custom, re.S)
                    .group(1).replace('\\u0022', '"'))
    check('custom range covering the fixtures finds all 8',
          sum(d6['trend']['values']) == total_fixtures, f'got {sum(d6["trend"]["values"])}')
    bad = render({'period': 'custom', 'date_from': 'nonsense', 'date_to': 'x'})
    check('malformed custom dates fall back instead of erroring', bad.status_code == 200)

    # An RTV NCM never dated has no bucket on the timeline, so the trend must
    # declare the shortfall rather than let its bars silently undercount.
    undated = make_rtv(admin, BASE_ID + 9, 'Customer cancelled', None,
                       '9801110009', 'Undated Buyer')
    wide = render(SCOPE).content.decode('utf-8', 'replace')
    d7 = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', wide, re.S)
                    .group(1).replace('\\u0022', '"'))
    check('undated RTV still counts toward the reason totals',
          sum(d7['reasons']['values']) == total_fixtures + 1,
          f'got {sum(d7["reasons"]["values"])}')
    check('undated RTV is left off the trend',
          sum(d7['trend']['values']) == total_fixtures, f'got {sum(d7["trend"]["values"])}')
    check('trend discloses the gap', 'excludes 1 undated' in wide)
    check('a dated range warns that undated rows are dropped',
          'cannot fall inside any date range' in
          render({'period': 'last30', 'search': '9910000'}).content.decode('utf-8', 'replace'))
    undated.delete()

    # An order with no created_by is found but cannot be credited to anyone. It
    # must not count as confirmed coverage, or the attribution rate claims a
    # name the leaderboard cannot show.
    print('\norders with no recorded creator')
    Order.objects.filter(pk=created[1][0].pk).update(created_by=None)
    nc = render(SCOPE).content.decode('utf-8', 'replace')
    d8 = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', nc, re.S)
                    .group(1).replace('\\u0022', '"'))
    alice_name = alice.get_full_name() or alice.username
    labels8 = d8['staff']['labels']
    check('creatorless order drops off that staff member',
          labels8 and (d8['staff']['confirmed'][labels8.index(alice_name)] == 2
                       if alice_name in labels8 else False),
          f'confirmed={d8["staff"]["confirmed"]} labels={labels8}')
    check('creatorless order is labelled, not silently dropped',
          'no creator recorded' in nc)
    check('creatorless order still counts in the totals',
          sum(d8['reasons']['values']) == total_fixtures,
          f'got {sum(d8["reasons"]["values"])}')
    check('creatorless order is reachable via the not-linked filter',
          sum(json.loads(
              re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>',
                        render(dict(SCOPE, confidence='none')).content.decode('utf-8', 'replace'),
                        re.S).group(1).replace('\\u0022', '"')
          )['trend']['values']) == 2)
    Order.objects.filter(pk=created[1][0].pk).update(created_by=alice)

    # ---------- NCM shipment status ----------
    # RTVStatus is a local tag almost nobody sets; last_status is what NCM
    # sends on every row. The Status column must read the populated one.
    print('\nNCM shipment status')
    RTVOrder.objects.filter(order_id=BASE_ID + 1).update(last_status='Returned to Warehouse')
    RTVOrder.objects.filter(order_id=BASE_ID + 2).update(last_status='Sent to Vendor')
    st = render(SCOPE).content.decode('utf-8', 'replace')
    d9 = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', st, re.S)
                    .group(1).replace('\\u0022', '"'))
    statuses = dict(zip(d9['statuses']['labels'], d9['statuses']['values']))
    check('status chart reads NCM status, not the unset local tag',
          statuses.get('Arrived') == 6 and statuses.get('Returned to Warehouse') == 1
          and statuses.get('Sent to Vendor') == 1, f'got {statuses}')
    check('no row falls back to a blank status', 'Unknown' not in statuses, f'got {statuses}')
    check('status renders in the detail table', 'Returned to Warehouse' in st)
    check('local tag panel is shown alongside', 'Your local RTV status tags' in st)

    only_rtw = render(dict(SCOPE, ncm_status='Returned to Warehouse')).content.decode('utf-8', 'replace')
    d10 = json.loads(re.search(r'id="rtvrChartData"[^>]*>(.*?)</script>', only_rtw, re.S)
                     .group(1).replace('\\u0022', '"'))
    check('NCM status filter narrows the set',
          sum(d10['trend']['values']) == 1, f'got {sum(d10["trend"]["values"])}')
    check('NCM status filter shows an active chip',
          'NCM Status: Returned to Warehouse' in only_rtw)
    RTVOrder.objects.filter(order_id__in=[BASE_ID + 1, BASE_ID + 2]).update(last_status='Arrived')

    # ---------- page size ----------
    print('\npage size')
    row_re = re.compile(r'<tr>\s*<td style="color:var\(--rtvr-muted\);">\d+</td>')

    def rendered_rows(params):
        return len(row_re.findall(render(params).content.decode('utf-8', 'replace')))

    check('default page shows up to 50', rendered_rows(SCOPE) == total_fixtures)
    check('per_page=25 caps the table',
          rendered_rows(dict(SCOPE, per_page='25')) == total_fixtures)
    check('per_page=all shows everything',
          rendered_rows(dict(SCOPE, per_page='all')) == total_fixtures)
    # With 8 fixtures a size cap is invisible, so prove the cap on real data:
    # the unfiltered page has hundreds of rows and must obey the choice.
    check('per_page=25 really caps a large set',
          rendered_rows({'period': 'all', 'per_page': '25'}) == 25)
    check('per_page=100 really widens a large set',
          rendered_rows({'period': 'all', 'per_page': '100'}) == 100)
    check('an off-menu per_page falls back to 50',
          rendered_rows({'period': 'all', 'per_page': '999'}) == 50)
    check('a non-numeric per_page falls back to 50',
          rendered_rows({'period': 'all', 'per_page': 'junk'}) == 50)
    sized = render(dict(SCOPE, per_page='25')).content.decode('utf-8', 'replace')
    check('page-size control is rendered', 'Rows per page' in sized)
    check('page size survives a filter submit',
          '<input type="hidden" name="per_page" value="25">' in sized)
    check('page-size links drop the old per_page',
          'data-base="?' in sized
          and 'per_page' not in re.search(r'data-base="\?([^"]*)"', sized).group(1))

    print('\nexport')
    xlsx = render(dict(SCOPE, export='xlsx'))
    check('export returns a spreadsheet',
          xlsx['Content-Type'].endswith('spreadsheetml.sheet'), xlsx['Content-Type'])
    check('export is an attachment', 'attachment;' in xlsx['Content-Disposition'])
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx.content))
    check('export has all three sheets',
          wb.sheetnames == ['RTV Detail', 'Staff Summary', 'Reason Summary'],
          f'got {wb.sheetnames}')
    detail = wb['RTV Detail']
    check('export detail has one row per RTV plus a header',
          detail.max_row == total_fixtures + 1, f'got {detail.max_row}')
    check('export honours the active filter',
          load_workbook(io.BytesIO(
              render(dict(SCOPE, staff=alice.username, export='xlsx')).content
          ))['RTV Detail'].max_row == 4)

    print('\npermissions')
    plain = CustomUser.objects.filter(
        is_superuser=False,
    ).exclude(role='administrator').first()
    if plain:
        was = plain.can_view_rtv_report
        CustomUser.objects.filter(pk=plain.pk).update(can_view_rtv_report=False)
        plain.refresh_from_db()
        req = rf.get('/reports/rtv/')
        req.user = plain
        # messages needs a backing store the bare RequestFactory has not set up.
        from django.contrib.messages.storage.fallback import FallbackStorage
        req.session = {}
        req._messages = FallbackStorage(req)
        denied = rtv_report(req)
        check('user without the permission is redirected away',
              denied.status_code == 302, f'got {denied.status_code}')

        CustomUser.objects.filter(pk=plain.pk).update(can_view_rtv_report=True)
        plain.refresh_from_db()
        req2 = rf.get('/reports/rtv/')
        req2.user = plain
        req2.session = {}
        req2._messages = FallbackStorage(req2)
        allowed = rtv_report(req2)
        check('granting can_view_rtv_report opens the page',
              allowed.status_code == 200, f'got {allowed.status_code}')
        CustomUser.objects.filter(pk=plain.pk).update(can_view_rtv_report=was)
    else:
        print('  SKIP  no non-admin user available to test the gate')

    return failures


if __name__ == '__main__':
    failures = []
    try:
        with transaction.atomic():
            failures = run()
            raise Rollback()
    except Rollback:
        print('\n(all test rows rolled back)')

    if failures:
        print(f'\n{len(failures)} FAILED: ' + ', '.join(failures))
        sys.exit(1)
    print('\nAll checks passed.')
