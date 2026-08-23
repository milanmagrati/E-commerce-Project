"""Verify the Staff Orders Log Excel export on the Staff Performance page.

Checks that ?orders_export=xlsx returns a real .xlsx whose rows match the
filtered queryset (all of it, not just the paginated page), that the money and
date cells carry usable Excel types, and that the status filter narrows the
export the same way it narrows the on-screen table.

    python test_staff_orders_export.py
"""
import os
import sys
import io

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'myproject.settings')
django.setup()

from django.contrib.auth import get_user_model
from django.db.models import Min, Max
from django.test import RequestFactory
from django.utils import timezone
from openpyxl import load_workbook

from dashboard.models import Order
from dashboard.views import staff_performance_analytics

User = get_user_model()

PASS, FAIL = [], []


def check(label, condition, detail=''):
    (PASS if condition else FAIL).append(label)
    print(('  [PASS] ' if condition else '  [FAIL] ') + label + (f' -> {detail}' if detail else ''))


def call_export(user, **params):
    params['orders_export'] = 'xlsx'
    request = RequestFactory().get('/staff-performance/', params)
    request.user = user
    # The view writes filter state to the session; give it a real dict-like one.
    from django.contrib.sessions.backends.db import SessionStore
    request.session = SessionStore()
    return staff_performance_analytics(request)


def load(response):
    return load_workbook(io.BytesIO(response.content))


def main():
    admin = User.objects.filter(is_superuser=True, is_active=True).first()
    if not admin:
        print('No superuser found - cannot exercise the view. Aborting.')
        return 1

    # Pick a window that actually contains orders in whatever DB this runs
    # against (the local copy trails production), so the value checks below
    # exercise real rows instead of an empty month.
    today = timezone.localtime(timezone.now()).date()
    span = Order.objects.filter(is_deleted=False).aggregate(
        first=Min('created_at'), last=Max('created_at')
    )
    if span['first']:
        start = timezone.localtime(span['first']).date()
        end = min(timezone.localtime(span['last']).date(), today)
    else:
        start = end = today
    PERIOD = dict(period='custom', custom_start=start.isoformat(), custom_end=end.isoformat())
    print(f'Acting as: {admin.username}   Period: custom ({start} .. {end})\n')

    print('1) Export returns a downloadable workbook')
    resp = call_export(admin, staff_filter='all', **PERIOD)
    check('HTTP 200', resp.status_code == 200, resp.status_code)
    check(
        'xlsx content type',
        'spreadsheetml.sheet' in resp['Content-Type'],
        resp['Content-Type'],
    )
    check(
        'attachment filename',
        'attachment' in resp['Content-Disposition'] and '.xlsx' in resp['Content-Disposition'],
        resp['Content-Disposition'],
    )

    wb = load(resp)
    check(
        'four sheets',
        wb.sheetnames == ['Staff Orders', 'Order Items', 'Staff Summary', 'Report Info'],
        wb.sheetnames,
    )

    orders_ws = wb['Staff Orders']
    exported = orders_ws.max_row - 1  # minus header

    print('\n2) Exports the whole filtered set, not just the visible page')
    from datetime import datetime, time as _time
    start_dt = timezone.make_aware(datetime.combine(start, _time.min))
    end_dt = timezone.make_aware(datetime.combine(end, _time.max))
    expected = Order.objects.filter(
        created_at__gte=start_dt, created_at__lte=end_dt, is_deleted=False
    ).count()
    check('row count == queryset count', exported == expected, f'{exported} vs {expected}')
    check('more than one page worth (or DB is small)', exported > 15 or expected <= 15, exported)

    if exported == 0:
        print('\nNo orders in range - skipping value checks.')
        return 0 if not FAIL else 1

    print('\n3) Cell values and formats are usable in Excel')
    headers = [c.value for c in orders_ws[1]]
    row2 = {h: c for h, c in zip(headers, orders_ws[2])}
    check('Amount is numeric', isinstance(row2['Amount'].value, (int, float)), type(row2['Amount'].value).__name__)
    check('Amount uses money format', row2['Amount'].number_format == '#,##0.00', row2['Amount'].number_format)
    check('Created At is a datetime', hasattr(row2['Created At'].value, 'year'), type(row2['Created At'].value).__name__)
    check('Created At is tz-naive (Excel-safe)', row2['Created At'].value.tzinfo is None)
    check('header row frozen', orders_ws.freeze_panes == 'A2', orders_ws.freeze_panes)
    check('auto filter set', orders_ws.auto_filter.ref is not None, orders_ws.auto_filter.ref)

    print('\n4) Row content matches the database record')
    first_order = Order.objects.filter(
        created_at__gte=start_dt, created_at__lte=end_dt, is_deleted=False
    ).select_related('created_by').order_by('-created_at').first()
    check('first row is newest order', row2['Order #'].value == first_order.order_number,
          f"{row2['Order #'].value} vs {first_order.order_number}")
    check('amount matches DB', abs(row2['Amount'].value - float(first_order.total_amount)) < 0.01,
          f"{row2['Amount'].value} vs {first_order.total_amount}")
    check('customer matches DB', row2['Customer'].value == (first_order.customer_name or ''),
          row2['Customer'].value)
    expected_status = (first_order.order_status or first_order.status or '').title()
    check('status matches the on-screen pill logic', row2['Status'].value == expected_status,
          f"{row2['Status'].value} vs {expected_status}")

    local_created = timezone.localtime(first_order.created_at)
    check(
        'created at rendered in Nepal time',
        row2['Created At'].value.strftime('%Y-%m-%d %H:%M') == local_created.strftime('%Y-%m-%d %H:%M'),
        f"{row2['Created At'].value} vs {local_created:%Y-%m-%d %H:%M}",
    )

    print('\n5) Order Items sheet lines up with OrderItem rows')
    items_ws = wb['Order Items']
    exported_items = items_ws.max_row - 1
    from dashboard.models import OrderItem
    expected_items = OrderItem.objects.filter(
        order__created_at__gte=start_dt,
        order__created_at__lte=end_dt,
        order__is_deleted=False,
    ).count()
    check('item row count matches', exported_items == expected_items,
          f'{exported_items} vs {expected_items}')

    print('\n6) Staff Summary totals reconcile with the detail sheet')
    summary_ws = wb['Staff Summary']
    s_headers = [c.value for c in summary_ws[1]]
    oi, ri = s_headers.index('Orders'), s_headers.index('Revenue')
    summary_orders = sum(r[oi].value for r in summary_ws.iter_rows(min_row=2))
    summary_revenue = sum(r[ri].value for r in summary_ws.iter_rows(min_row=2))
    amount_idx = headers.index('Amount')
    detail_revenue = sum(r[amount_idx].value for r in orders_ws.iter_rows(min_row=2))
    check('summary order count == detail rows', summary_orders == exported,
          f'{summary_orders} vs {exported}')
    check('summary revenue == detail revenue', abs(summary_revenue - detail_revenue) < 1.0,
          f'{summary_revenue} vs {detail_revenue}')

    print('\n7) Status filter narrows the export the same way it narrows the table')
    a_status = (first_order.order_status or first_order.status or '')
    if a_status:
        resp2 = call_export(admin, staff_filter='all', orders_status=a_status, **PERIOD)
        ws2 = load(resp2)['Staff Orders']
        rows2 = ws2.max_row - 1
        from django.db.models import Q
        expected2 = Order.objects.filter(
            created_at__gte=start_dt, created_at__lte=end_dt, is_deleted=False
        ).filter(Q(status__iexact=a_status) | Q(order_status__iexact=a_status)).count()
        check(f'status="{a_status}" row count matches', rows2 == expected2, f'{rows2} vs {expected2}')
        check('filtered export is a subset', rows2 <= exported, f'{rows2} <= {exported}')
        info = {r[0].value: r[1].value for r in load(resp2)['Report Info'].iter_rows(min_row=2)}
        check('Report Info records the status filter',
              str(info.get('Status Filter', '')).lower() == a_status.lower(),
              info.get('Status Filter'))

    print('\n8) Per-staff filter is honoured')
    staff_with_orders = Order.objects.filter(
        created_at__gte=start_dt, created_at__lte=end_dt, is_deleted=False,
        created_by__isnull=False,
    ).values_list('created_by_id', flat=True).first()
    if staff_with_orders:
        resp3 = call_export(admin, staff_filter=str(staff_with_orders), **PERIOD)
        wb3 = load(resp3)
        rows3 = wb3['Staff Orders'].max_row - 1
        expected3 = Order.objects.filter(
            created_at__gte=start_dt, created_at__lte=end_dt, is_deleted=False,
            created_by_id=staff_with_orders,
        ).count()
        check('staff-filtered row count matches', rows3 == expected3, f'{rows3} vs {expected3}')
        names = {r[0].value for r in wb3['Staff Summary'].iter_rows(min_row=2)}
        check('only one staff in summary', len(names) == 1, names)

    print(f'\n{"=" * 60}\n{len(PASS)} passed, {len(FAIL)} failed')
    for f in FAIL:
        print('  FAILED: ' + f)
    return 1 if FAIL else 0


if __name__ == '__main__':
    sys.exit(main())
